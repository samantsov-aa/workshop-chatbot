"""Builds worksheet PDFs and the slide deck (PDF + PPTX) into the workshop folder."""
import json
import re
import shutil
import subprocess
import sys
from pathlib import Path

import markdown
import pypdfium2 as pdfium
from pptx import Presentation
from pptx.util import Emu

DECK = Path(sys.argv[1]).resolve()  # folder with deck.json and slides/
OUT = Path(sys.argv[2]).resolve()   # workshop folder (worksheets/ and slides/ inside)
BROWSERS = [r"C:\Program Files\Google\Chrome\Application\chrome.exe",
            r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe"]
EDGE = next((b for b in BROWSERS if Path(b).exists()), None) or sys.exit("Need Chrome or Edge")
WORK = OUT / ".build_tmp"
shutil.rmtree(WORK, ignore_errors=True)
WORK.mkdir(exist_ok=True)


def print_pdf(src: Path, out: Path):
    """Headless Edge sometimes exits without writing anything, so retry until the file appears."""
    out.unlink(missing_ok=True)
    for attempt in range(5):
        # a fresh profile per run: a lingering Edge process holding the old one makes the next run a no-op
        profile = WORK / f"profile_{out.stem}_{attempt}"
        subprocess.run([EDGE, "--headless=new", "--disable-gpu", "--hide-scrollbars",
                        f"--user-data-dir={profile}", "--no-first-run",
                        "--virtual-time-budget=8000", "--no-pdf-header-footer",
                        f"--print-to-pdf={out}", src.as_uri()],
                       stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        if out.exists():
            return
    sys.exit(f"Edge failed to print {out.name}")


# ---------- worksheets ----------
SHEET_CSS = """
@import url('https://fonts.googleapis.com/css2?family=Nunito:wght@400;700;800&family=Rubik:wght@500;600&family=JetBrains+Mono&display=swap');
@page { size: A4; margin: 16mm 15mm; }
body { font-family: 'Nunito', 'Segoe UI', sans-serif; font-size: 11pt; color: #1B2340; line-height: 1.45; }
h1, h2 { font-family: 'Rubik', 'Segoe UI', sans-serif; font-weight: 600; }
h1 { font-size: 20pt; margin: 0 0 8pt; }
h2 { font-size: 14pt; margin: 16pt 0 6pt; color: #9A4512; page-break-after: avoid; }
table { border-collapse: collapse; width: 100%; margin: 6pt 0; page-break-inside: avoid; }
th, td { border: 1px solid #C9CED9; padding: 5pt 7pt; text-align: left; vertical-align: top; }
th { background: #EAF0FB; }
td:empty::after { content: "\\00a0"; }
tr { height: 22pt; }
code { font-family: 'JetBrains Mono', Consolas, monospace; font-size: 9.5pt; background: #F1EFE8; padding: 0 3pt; border-radius: 3pt; }
ul, ol { padding-left: 18pt; }
li { margin: 2pt 0; }
a { color: #3B5BA9; }
"""

for md_file in sorted((OUT / "worksheets").glob("*.md")):
    html_body = markdown.markdown(md_file.read_text(encoding="utf-8"), extensions=["tables"])
    html_body = html_body.replace("<li>[ ] ", "<li style='list-style:none;margin-left:-14pt'>☐ ")
    page = WORK / (md_file.stem + ".html")
    page.write_text(f"<!doctype html><html><head><meta charset='utf-8'><style>{SHEET_CSS}</style>"
                    f"</head><body>{html_body}</body></html>", encoding="utf-8")
    print_pdf(page, md_file.with_suffix(".pdf"))
    print("pdf:", md_file.with_suffix(".pdf").name)

# ---------- slides ----------
deck = json.loads((DECK / "deck.json").read_text(encoding="utf-8"))
fonts = "".join(f"<link rel='stylesheet' href='{f['href']}'>" for f in deck["faces"].values())

EMOJI = {"Chat": "💬", "Book": "📖", "Search": "🔍", "Lightbulb": "💡", "Settings": "⚙️",
         "Key": "🔑", "Warning": "⚠️", "Tool": "✂️", "Database": "🗄️", "CheckCircle": "✅",
         "Lock": "🔒", "GraduationCap": "🎓", "Star": "⭐", "Globe": "🌐"}
ARROW = "clip-path:polygon(0 30%,60% 30%,60% 0,100% 50%,60% 100%,60% 70%,0 70%)"
SLIDE_CSS = """
* { margin: 0; box-sizing: border-box; }
@page { size: 1920px 1080px; margin: 0; }
body { width: 1920px; }
section { width: 1920px; height: 1080px; position: relative; overflow: hidden; page-break-after: always; }
ul { padding-left: 1.2em; }
table { border-collapse: collapse; }
th, td { padding: 0.35em 0.6em; border-bottom: 2px solid #D5D9E3; }
th { border-bottom: 3px solid #1B2340; }
"""


def convert(html: str) -> tuple[str, str]:
    notes = ""
    m = re.search(r"<aside>(.*?)</aside>", html, re.S)
    if m:
        notes = m.group(1).strip()
        html = html.replace(m.group(0), "")

    def icon(m):
        name, style = m.group(1), m.group(2)
        size = int(re.search(r"height:(\d+)px", style).group(1))
        return (f"<span style=\"{style};display:inline-flex;align-items:center;justify-content:center;"
                f"flex:none;font-size:{int(size * 0.82)}px;line-height:1;font-family:'Segoe UI Emoji'\">"
                f"{EMOJI[name]}</span>")

    html = re.sub(r'<x-icon name="(\w+)" style="([^"]*)"></x-icon>', icon, html)

    def shape(m):
        kind, style = m.group(1), m.group(2)
        extra = "border-radius:50%" if kind == "ellipse" else ARROW
        return f'<div style="{style};flex:none;{extra}"></div>'

    html = re.sub(r'<x-shape kind="([\w-]+)" style="([^"]*)"></x-shape>', shape, html)
    return html, notes


sections, notes = [], []
for sid in deck["order"]:
    html, note = convert((DECK / "slides" / f"{sid}.html").read_text(encoding="utf-8"))
    sections.append(html)
    notes.append(note)


def page(body):
    return (f"<!doctype html><html><head><meta charset='utf-8'>{fonts}<style>{SLIDE_CSS}</style>"
            f"</head><body>{body}</body></html>")


slides_out = OUT / "slides"
slides_out.mkdir(exist_ok=True)
full = WORK / "deck.html"
full.write_text(page("".join(sections)), encoding="utf-8")
print_pdf(full, slides_out / "Презентація.pdf")
print("pdf: Презентація.pdf")

pdf = pdfium.PdfDocument(slides_out / "Презентація.pdf")
assert len(pdf) == len(sections), f"{len(pdf)} pages for {len(sections)} slides"
pngs = []
for i, pg in enumerate(pdf, 1):
    png = WORK / f"slide{i:02d}.png"
    pg.render(scale=1920 / pg.get_width()).to_pil().save(png)
    pngs.append(png)
pdf.close()

prs = Presentation()
prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
for png, note in zip(pngs, notes):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.shapes.add_picture(str(png), 0, 0, prs.slide_width, prs.slide_height)
    s.notes_slide.notes_text_frame.text = note
prs.save(slides_out / "Презентація.pptx")
print("pptx: Презентація.pptx")
shutil.rmtree(WORK, ignore_errors=True)
